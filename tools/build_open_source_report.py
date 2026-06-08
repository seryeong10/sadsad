from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "오픈소스_활용_보고서_A-VATA.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def style_paragraph(paragraph, size=10.5, bold=False, color="000000"):
    for run in paragraph.runs:
        run.font.name = "Malgun Gothic"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.12


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    run = p.add_run(text)
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.bold = True
    run.font.color.rgb = RGBColor(33, 91, 151)
    run.font.size = Pt(16 if level == 1 else 13)
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_body(doc, text, size=10.5):
    p = doc.add_paragraph(text)
    style_paragraph(p, size=size)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)
        style_paragraph(p, size=10.3)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        header_cells[idx].text = header
        header_cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(header_cells[idx], "F2F4F7")
        set_cell_margins(header_cells[idx])
        for p in header_cells[idx].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            style_paragraph(p, size=9.3, bold=True)

    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cells[idx])
            for p in cells[idx].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                style_paragraph(p, size=8.7)
    set_table_width(table, widths)
    doc.add_paragraph()
    return table


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    styles = doc.styles
    styles["Normal"].font.name = "Malgun Gothic"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("오픈소스 활용 보고서 (캡스톤 연계)")
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(23)
    run.font.bold = True
    run.font.color.rgb = RGBColor(23, 55, 94)
    title.paragraph_format.space_after = Pt(14)

    add_heading(doc, "1. 프로젝트 기본 정보", 1)
    add_body(doc, "프로젝트명: A-VATA")
    add_body(doc, "팀명 / 팀원: 룩킹포유 / 박세령, 김용준, 윤선경, 정미영")
    add_body(doc, "GitHub 저장소: https://github.com/seryeong10/capstone_avata")
    add_body(
        doc,
        "프로젝트 개요: A-VATA는 사용자의 정면 사진을 기반으로 전신 아바타의 얼굴과 헤어를 편집하고, 네이버 쇼핑 API로 수집한 의류 데이터를 활용해 향후 사용자 무드 기반 코디 추천으로 확장하는 Flutter + FastAPI 기반 아바타/패션 추천 서비스이다.",
    )

    add_heading(doc, "2. 오픈소스 사용 목록", 1)
    add_table(
        doc,
        ["오픈소스명", "버전", "사용 목적", "적용 기능", "라이선스"],
        [
            ["Flutter", "SDK 3.4 이상", "앱 프론트엔드 개발", "성별 선택, 사진 업로드, 아바타 결과 화면, 홈/무드 화면 구성", "BSD-3-Clause"],
            ["FastAPI", "0.115.12", "백엔드 API 서버 구축", "사용자 사진과 기본 아바타 이미지를 받는 /generate-avatar API 구현", "MIT"],
            ["OpenAI Python SDK", "버전 명시 예정", "OpenAI 이미지 편집 API 연동", "아바타 얼굴/헤어 편집 요청 및 base64 PNG 결과 수신", "Apache-2.0"],
            ["Pillow", "11.2.1", "이미지 후처리", "생성 이미지의 밝은 배경 영역 투명화", "HPND"],
            ["Uvicorn", "0.34.2", "FastAPI 실행 서버", "로컬 개발 서버 실행 및 API 테스트", "BSD-3-Clause"],
            ["requests", "2.32.3", "네이버 쇼핑 API 호출 및 의류 데이터 수집", "검색어별 상품 데이터 요청 및 JSON 응답 처리", "Apache-2.0"],
            ["pandas", "버전 명시 예정", "수집 데이터 정리 및 CSV 저장", "의류 상품 데이터를 DataFrame으로 정리하고 CSV로 저장", "BSD-3-Clause"],
            ["python-dotenv", "버전 명시 예정", "API 키 및 환경변수 관리", "NAVER_CLIENT_ID, NAVER_CLIENT_SECRET, OPENAI_API_KEY 로드", "BSD-3-Clause"],
        ],
        [1.12, 0.78, 1.42, 2.05, 0.95],
    )

    add_heading(doc, "3. 오픈소스 선택 이유", 1)
    add_bullets(
        doc,
        [
            "Flutter는 하나의 코드베이스로 Android, iOS, 데스크톱, 웹 UI를 구성할 수 있어 캡스톤 MVP를 빠르게 구현하기에 적합하다.",
            "FastAPI는 비동기 파일 업로드와 API 문서화가 간단해 이미지 생성 서버를 짧은 코드로 구성할 수 있다.",
            "OpenAI Python SDK는 이미지 편집 API 호출을 표준화해 아바타 생성 기능을 빠르게 연결할 수 있다.",
            "Pillow는 생성된 PNG의 가장자리 배경을 픽셀 단위로 확인하고 투명화하는 이미지 후처리에 적합하다.",
            "requests와 pandas는 네이버 쇼핑 API 호출, 상품 데이터 정리, CSV 저장으로 이어지는 데이터 수집 파이프라인을 간단하게 구현할 수 있다.",
            "python-dotenv는 API 인증키를 코드에 직접 작성하지 않고 환경변수로 관리할 수 있어 보안과 유지보수에 유리하다.",
        ],
    )

    add_heading(doc, "4. 오픈소스 적용 상세", 1)
    add_table(
        doc,
        ["오픈소스명", "사용 위치\n(모듈/기능)", "적용 방식\n(사용/수정/확장)", "코드 설명"],
        [
            ["Flutter", "frontend/lib/pages, widgets", "사용", "성별 선택, 사진 업로드, 로딩, 결과, 아바타 홈 화면을 위젯 기반으로 구성했다."],
            ["FastAPI", "backend/app/main.py", "사용/확장", "POST /generate-avatar 엔드포인트에서 업로드 파일을 읽고 이미지 편집 요청을 수행한다."],
            ["OpenAI Python SDK", "backend/app/main.py", "사용", "OpenAI 클라이언트로 Images edit API를 호출하고 base64 PNG 결과를 수신한다."],
            ["Pillow", "backend/app/main.py", "확장", "_remove_light_edge_background 함수에서 RGBA 픽셀을 탐색해 밝은 가장자리 배경을 투명화한다."],
            ["Uvicorn", "backend 실행 환경", "사용", "uvicorn app.main:app 명령으로 로컬 백엔드 서버를 실행한다."],
            ["requests", "backend/app/naver_shop_test.py", "사용", "네이버 쇼핑 검색 API에 query와 display 파라미터를 보내 의류 상품 JSON 데이터를 수집한다."],
            ["pandas", "backend/app/naver_shop_test.py", "사용", "수집한 상품명, 이미지 URL, 링크, 가격, 쇼핑몰명, 카테고리 정보를 DataFrame으로 만들고 CSV로 저장한다."],
            ["python-dotenv", "backend/app/main.py, naver_shop_test.py", "사용", ".env 파일에서 OpenAI 및 네이버 API 인증 정보를 불러와 코드와 키를 분리한다."],
        ],
        [1.08, 1.55, 1.25, 2.62],
    )

    doc.add_section(WD_SECTION_START.NEW_PAGE)
    add_heading(doc, "5. 오픈소스 수정 및 확장 내용", 1)
    add_body(
        doc,
        "오픈소스 자체의 소스코드를 직접 수정하지는 않았고, 프로젝트 요구사항에 맞게 조합 및 확장하여 사용했다. 특히 Flutter 화면 흐름, FastAPI 파일 업로드 API, OpenAI 이미지 편집 연동, Pillow 기반 이미지 후처리, requests/pandas 기반 네이버 쇼핑 데이터 수집을 연결해 A-VATA의 핵심 기능을 구성했다.",
    )
    add_bullets(
        doc,
        [
            "Pillow 확장 활용: 생성 이미지의 네 모서리에서 시작해 밝고 중립적인 배경 픽셀을 탐색하고 alpha 값을 0으로 바꾸는 flood-fill 방식 후처리를 구현했다.",
            "Flutter 앱/서버 연동: 기본 아바타 이미지와 사용자 얼굴 사진을 함께 서버로 전송하고, 반환된 base64 PNG를 앱 화면에 표시하도록 구성했다.",
            "FastAPI 엔드포인트 확장: base_avatar와 user_face 두 파일을 동시에 받아 OpenAI 이미지 편집 요청의 입력 파일 목록으로 전달하도록 구현했다.",
            "의류 데이터 구조화: 네이버 쇼핑 API로 수집한 약 200개 상품 데이터를 CSV로 저장하고, category_simple, situation_tags, mood_tags, style_tags, season_tags 컬럼을 추가해 추천 기능에 활용하기 쉬운 형태로 구조화했다.",
        ],
    )

    add_heading(doc, "6. 라이선스 분석", 1)
    add_table(
        doc,
        ["오픈소스명", "라이선스 유형", "상업적 사용 가능 여부", "주의사항"],
        [
            ["Flutter", "BSD-3-Clause", "가능", "저작권 고지와 라이선스 문구를 유지해야 한다."],
            ["FastAPI", "MIT", "가능", "배포 시 라이선스 및 저작권 고지를 포함하는 것이 필요하다."],
            ["Uvicorn", "BSD-3-Clause", "가능", "소스 또는 바이너리 배포 시 BSD 고지 조건을 확인해야 한다."],
            ["OpenAI Python SDK", "Apache-2.0", "가능", "SDK 라이선스와 별도로 OpenAI API 이용 약관 및 비용 정책을 확인해야 한다."],
            ["Pillow", "HPND", "가능", "Pillow 라이선스 전문과 저작권 고지를 함께 보관해야 한다."],
            ["requests", "Apache-2.0", "가능", "NOTICE 파일 및 특허 조항 등 Apache 2.0 조건을 확인해야 한다."],
            ["pandas", "BSD-3-Clause", "가능", "저작권 고지와 라이선스 문구를 유지해야 한다."],
            ["python-dotenv", "BSD-3-Clause", "가능", "저작권 고지와 라이선스 문구를 유지해야 한다."],
        ],
        [1.42, 1.12, 1.4, 2.35],
    )

    add_heading(doc, "7. 오픈소스 활용 중 문제 및 해결", 1)
    add_body(doc, "문제 상황:")
    add_bullets(
        doc,
        [
            "Flutter 에셋 이미지는 그대로 multipart 파일로 보낼 수 없어 서버가 요구하는 File 경로로 변환하는 과정이 필요했다.",
            "Android 에뮬레이터에서는 로컬 서버 주소로 127.0.0.1을 쓰면 앱 내부 에뮬레이터 자신을 가리켜 백엔드 연결이 실패할 수 있었다.",
            "생성 이미지의 배경이 완전히 투명하지 않은 경우 앱 화면에서 밝은 테두리나 배경이 남을 수 있었다.",
            "네이버 쇼핑 API 응답에는 HTML 강조 태그가 포함되어 상품명을 그대로 저장하면 데이터 정리가 어려웠다.",
            "OpenAI Python SDK, pandas, python-dotenv는 코드에서 사용되지만 현재 requirements.txt에는 버전이 명시되어 있지 않아 새 환경 설치 시 보완이 필요하다.",
        ],
    )
    add_body(doc, "해결 방법:")
    add_bullets(
        doc,
        [
            "path_provider의 임시 디렉터리와 rootBundle.load를 사용해 기본 아바타 에셋을 PNG 파일로 저장한 뒤 http.MultipartFile.fromPath로 전송했다.",
            "Android 에뮬레이터 기준 API 주소를 http://10.0.2.2:8000/generate-avatar로 설정하고 README에 실행 환경별 주소를 정리했다.",
            "Pillow로 RGBA 이미지를 열고 가장자리와 연결된 밝은 중립색 픽셀만 투명 처리해 아바타 본체 손상을 줄였다.",
            "상품명 저장 전에 <b>, </b> 태그를 제거하고, pandas DataFrame으로 컬럼을 맞춘 뒤 CSV 파일로 저장했다.",
            "배포 전 requirements.txt에 openai, pandas, python-dotenv 버전을 추가하고, README의 환경변수 설정 절차를 따라 실행하도록 정리할 필요가 있다.",
        ],
    )

    add_heading(doc, "8. 오픈소스 활용 평가", 1)
    add_body(doc, "장점:")
    add_bullets(
        doc,
        [
            "Flutter와 FastAPI 조합 덕분에 모바일 UI와 이미지 생성 API를 분리해 빠르게 MVP를 만들 수 있었다.",
            "OpenAI Python SDK와 Pillow를 함께 사용해 이미지 생성 요청부터 결과 이미지 후처리까지 하나의 파이프라인으로 구성할 수 있었다.",
            "requests와 pandas를 활용해 네이버 쇼핑 API 기반 의류 상품 데이터를 빠르게 수집하고 CSV 데이터셋으로 구조화할 수 있었다.",
            "각 라이브러리의 문서와 커뮤니티가 풍부해 오류 원인 파악과 기능 확장이 비교적 쉬웠다.",
        ],
    )
    add_body(doc, "단점:")
    add_bullets(
        doc,
        [
            "Flutter 패키지는 플랫폼별 권한 설정과 에뮬레이터 네트워크 주소 차이를 고려해야 해 초기 설정 비용이 있다.",
            "이미지 생성 결과 품질은 외부 API와 프롬프트에 영향을 받기 때문에 동일한 입력에서도 결과 편차가 생길 수 있다.",
            "네이버 쇼핑 API와 OpenAI API는 오픈소스가 아닌 외부 API이므로 이용 조건, 호출 제한, 비용 정책을 별도로 확인해야 한다.",
            "여러 오픈소스를 함께 사용하므로 requirements.txt와 pubspec.yaml의 버전, 라이선스, 보안 업데이트를 지속적으로 관리해야 한다.",
        ],
    )

    add_heading(doc, "9. 현재 구현 범위 및 향후 확장", 1)
    add_body(doc, "현재 구현 완료 범위:")
    add_bullets(
        doc,
        [
            "Flutter + FastAPI 기반 앱/서버 연결",
            "사용자 사진 업로드 및 기본 아바타 이미지 전송",
            "OpenAI 이미지 편집 API를 활용한 아바타 얼굴/헤어 편집",
            "Pillow를 활용한 생성 이미지 배경 투명화",
            "네이버 쇼핑 API 기반 의류 데이터 수집",
            "상품명, 이미지 URL, 링크, 가격, 쇼핑몰명, 카테고리 정보 CSV 저장",
            "상황/무드/스타일/계절 태그 기반 의류 데이터 구조화",
        ],
    )
    add_body(doc, "향후 확장 예정:")
    add_bullets(
        doc,
        [
            "의류 데이터 전처리 고도화",
            "상품 메타데이터 DB 저장",
            "의류 설명 및 스타일 태그 임베딩",
            "ChromaDB 또는 FAISS 기반 벡터 DB 저장",
            "사용자 무드/아바타 기반 유사도 의상 추천",
            "AI 피팅 및 코디 추천 기능 고도화",
        ],
    )

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("A-VATA 오픈소스 활용 보고서")
    style_paragraph(footer, size=9, color="666666")

    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
